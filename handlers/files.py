import io
import json
import re

from telegram import Update
from telegram.constants import ChatAction
from telegram.ext import ContextTypes

import ai_client
import db
import utils
from handlers.keyboards import build_main_menu

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Pt
from fpdf import FPDF

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    HAS_RTL = True
except Exception:
    HAS_RTL = False

KIND_LABELS = {"pdf": "PDF", "docx": "Word", "xlsx": "Excel", "pptx": "PowerPoint", "txt": "متنی"}


def _rtl_line(line: str) -> str:
    if HAS_RTL and utils.contains_persian(line):
        return get_display(arabic_reshaper.reshape(line))
    return line


# ------------------------------------------------------------------ PDF

def build_pdf(title: str, body: str) -> bytes:
    font_path = utils.ensure_persian_font()
    pdf = FPDF()
    pdf.add_page()
    if font_path:
        pdf.add_font("Persian", "", font_path, uni=True)
        pdf.set_font("Persian", size=16)
    else:
        pdf.set_font("Helvetica", size=16)

    pdf.multi_cell(0, 10, _rtl_line(title), align="C")
    pdf.ln(4)
    pdf.set_font_size(12)
    for paragraph in body.split("\n"):
        align = "R" if utils.contains_persian(paragraph) else "L"
        pdf.multi_cell(0, 8, _rtl_line(paragraph), align=align)
    return bytes(pdf.output())


# ------------------------------------------------------------------ DOCX

def build_docx(title: str, body: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for paragraph in body.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ------------------------------------------------------------------ entry point used by the menu

async def generate_file(update: Update, context: ContextTypes.DEFAULT_TYPE, kind: str, topic: str):
    user_id = update.effective_user.id
    topic = (topic or "").strip()
    if not topic:
        await update.message.reply_text(
            "متن نمی‌تواند خالی باشد. دوباره از دکمه «📄 ساخت فایل» استفاده کن.",
            reply_markup=build_main_menu(user_id),
        )
        return

    if kind == "xlsx":
        await _generate_xlsx(update, context, topic)
    elif kind == "pptx":
        await _generate_pptx(update, context, topic)
    elif kind == "txt":
        await _generate_txt(update, context, topic)
    elif kind in ("pdf", "docx"):
        await _generate_text_document(update, context, kind, topic)
    else:
        await update.message.reply_text("نوع فایل ناشناخته است.", reply_markup=build_main_menu(user_id))


# ------------------------------------------------------------------ pdf / docx (shared flow)

async def _generate_text_document(update, context, kind, prompt):
    user_id = update.effective_user.id
    raw_mode = prompt.startswith("raw:")
    status = await update.message.reply_text(f"⏳ در حال آماده‌سازی فایل {KIND_LABELS[kind]}...")
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action=ChatAction.UPLOAD_DOCUMENT)

    if raw_mode:
        title = "Document"
        body = prompt[4:].strip()
    else:
        try:
            content = ai_client.chat_completion([
                {"role": "system", "content": "You write clean, well structured document content. "
                                               "Respond in the user's language. Output plain text only, "
                                               "no markdown symbols. First line = a short title."},
                {"role": "user", "content": f"Write the content for a {kind} document about: {prompt}"},
            ], model=db.get_user_model(user_id))
        except ai_client.AnyModelError as e:
            await status.edit_text(f"❌ خطا در تولید محتوا:\n{e}")
            await update.message.reply_text("منوی اصلی:", reply_markup=build_main_menu(user_id))
            return
        lines = content.strip().split("\n", 1)
        title = lines[0][:120]
        body = lines[1].strip() if len(lines) > 1 else content

    filename = utils.safe_filename(title) + f".{kind}"

    try:
        file_bytes = build_pdf(title, body) if kind == "pdf" else build_docx(title, body)
    except Exception as e:
        await status.edit_text(f"❌ خطا در ساخت فایل:\n{e}")
        await update.message.reply_text("منوی اصلی:", reply_markup=build_main_menu(user_id))
        return

    await status.delete()
    bio = io.BytesIO(file_bytes)
    bio.name = filename
    await update.message.reply_document(document=bio, filename=filename, reply_markup=build_main_menu(user_id))
    db.bump_counter(user_id, "file_count")


# ------------------------------------------------------------------ xlsx

async def _generate_xlsx(update, context, prompt):
    user_id = update.effective_user.id
    status = await update.message.reply_text("⏳ در حال ساخت اکسل...")
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action=ChatAction.UPLOAD_DOCUMENT)

    try:
        content = ai_client.chat_completion([
            {"role": "system", "content": (
                "You output ONLY raw CSV data (comma separated), nothing else - no markdown, no code fences, "
                "no explanations. The first row must be column headers. Respond in the user's language."
            )},
            {"role": "user", "content": f"Create a CSV table about: {prompt}"},
        ], model=db.get_user_model(user_id))
    except ai_client.AnyModelError as e:
        await status.edit_text(f"❌ خطا در تولید داده:\n{e}")
        await update.message.reply_text("منوی اصلی:", reply_markup=build_main_menu(user_id))
        return

    rows = [r for r in content.strip().split("\n") if r.strip()]
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.rightToLeft = utils.contains_persian(content)
    for r in rows:
        cells = [c.strip().strip('"') for c in re.split(r",(?=(?:[^\"]*\"[^\"]*\")*[^\"]*$)", r)]
        ws.append(cells)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 50)

    bio = io.BytesIO()
    wb.save(bio)
    filename = utils.safe_filename(prompt.split()[0] if prompt else "table") + ".xlsx"
    bio.seek(0)
    bio.name = filename

    await status.delete()
    await update.message.reply_document(document=bio, filename=filename, reply_markup=build_main_menu(user_id))
    db.bump_counter(user_id, "file_count")


# ------------------------------------------------------------------ pptx

async def _generate_pptx(update, context, prompt):
    user_id = update.effective_user.id
    status = await update.message.reply_text("⏳ در حال ساخت پاورپوینت...")
    await context.bot.send_chat_action(chat_id=update.effective_chat.id, action=ChatAction.UPLOAD_DOCUMENT)

    schema_hint = (
        'Respond ONLY with valid JSON, no markdown fences, in this exact shape: '
        '{"title": "...", "slides": [{"title": "...", "bullets": ["...", "..."]}]}. '
        "5 to 8 slides. Respond in the user's language."
    )
    try:
        content = ai_client.chat_completion([
            {"role": "system", "content": schema_hint},
            {"role": "user", "content": f"Create a slide outline about: {prompt}"},
        ], model=db.get_user_model(user_id))
        content = re.sub(r"^```(json)?|```$", "", content.strip(), flags=re.MULTILINE).strip()
        outline = json.loads(content)
    except Exception as e:
        await status.edit_text(f"❌ خطا در تولید محتوای اسلایدها:\n{e}")
        await update.message.reply_text("منوی اصلی:", reply_markup=build_main_menu(user_id))
        return

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = outline.get("title", prompt)[:120]
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "ساخته‌شده توسط ربات هوش مصنوعی"

    for slide_data in outline.get("slides", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(slide_data.get("title", ""))[:120]
        body = slide.placeholders[1].text_frame
        bullets = slide_data.get("bullets", [])
        for i, bullet in enumerate(bullets):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = str(bullet)
            p.font.size = Pt(18)

    bio = io.BytesIO()
    prs.save(bio)
    filename = utils.safe_filename(outline.get("title", prompt)) + ".pptx"
    bio.seek(0)
    bio.name = filename

    await status.delete()
    await update.message.reply_document(document=bio, filename=filename, reply_markup=build_main_menu(user_id))
    db.bump_counter(user_id, "file_count")


# ------------------------------------------------------------------ txt

async def _generate_txt(update, context, prompt):
    user_id = update.effective_user.id
    status = await update.message.reply_text("⏳ در حال آماده‌سازی...")
    try:
        content = ai_client.chat_completion([
            {"role": "system", "content": "Write plain text content. Respond in the user's language."},
            {"role": "user", "content": prompt},
        ], model=db.get_user_model(user_id))
    except ai_client.AnyModelError as e:
        await status.edit_text(f"❌ خطا:\n{e}")
        await update.message.reply_text("منوی اصلی:", reply_markup=build_main_menu(user_id))
        return
    bio = io.BytesIO(content.encode("utf-8"))
    filename = utils.safe_filename(prompt.split()[0] if prompt else "note") + ".txt"
    bio.name = filename
    await status.delete()
    await update.message.reply_document(document=bio, filename=filename, reply_markup=build_main_menu(user_id))
    db.bump_counter(user_id, "file_count")
